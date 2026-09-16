from __future__ import annotations

import html
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse

from app.api.dependencies.auth import require_super_admin, get_current_user
from app.api.dependencies.database import require_database
from app.repositories import essential8 as essential8_repo
from app.repositories import user_companies as user_company_repo
from app.schemas.essential8 import (
    ApprovalStatus,
    CompanyEssential8AuditResponse,
    CompanyEssential8ComplianceCreate,
    CompanyEssential8ComplianceResponse,
    CompanyEssential8ComplianceSummary,
    CompanyEssential8ComplianceUpdate,
    CompanyEssential8RequirementBulkUpdate,
    CompanyEssential8RequirementComplianceCreate,
    CompanyEssential8RequirementComplianceResponse,
    CompanyEssential8RequirementComplianceUpdate,
    ComplianceStatus,
    Essential8RequirementEvidenceResponse,
    Essential8ControlResponse,
    Essential8ControlWithRequirementsResponse,
    Essential8RequirementResponse,
)

router = APIRouter(prefix="/api/essential8", tags=["Essential 8 Compliance"])


async def _get_company_membership(user: dict, company_id: int) -> dict | None:
    user_id = user.get("id")
    if user_id is None:
        return None
    return await user_company_repo.get_user_company(user_id, company_id)


async def _assert_company_compliance_access(user: dict, company_id: int, *, write: bool = False) -> dict | None:
    is_super_admin = bool(user.get("is_super_admin", False))
    user_company_id = user.get("company_id")
    membership = await _get_company_membership(user, company_id)
    if not is_super_admin and user_company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only access compliance for your own company",
        )
    if is_super_admin:
        return membership
    if not membership or not membership.get("can_view_compliance"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have access to compliance requirements for this company",
        )
    if write and not membership.get("is_admin"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Administrator access is required to manage compliance requirements",
        )
    return membership


def _requirement_upload_dir() -> Path:
    from app import main as main_module

    path = main_module._private_uploads_path / "compliance" / "essential8"
    path.mkdir(parents=True, exist_ok=True)
    return path


@router.get("/controls", response_model=list[Essential8ControlResponse])
async def list_controls(
    _: None = Depends(require_database),
    __: dict = Depends(get_current_user),
):
    """List all Essential 8 controls"""
    controls = await essential8_repo.list_essential8_controls()
    return controls


@router.get("/controls/{control_id}", response_model=Essential8ControlResponse)
async def get_control(
    control_id: int,
    _: None = Depends(require_database),
    __: dict = Depends(get_current_user),
):
    """Get a specific Essential 8 control"""
    control = await essential8_repo.get_essential8_control(control_id)
    if not control:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Control not found",
        )
    return control


@router.get(
    "/companies/{company_id}/compliance",
    response_model=list[CompanyEssential8ComplianceResponse],
)
async def list_company_compliance(
    company_id: int,
    status_filter: Optional[ComplianceStatus] = None,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """List compliance records for a company"""
    # Super admins can view any company, regular users can only view their own
    is_super_admin = user.get("is_super_admin", False)
    user_company_id = user.get("company_id")
    
    if not is_super_admin and user_company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only view compliance for your own company",
        )
    
    records = await essential8_repo.list_company_compliance(
        company_id=company_id,
        status=status_filter,
    )
    return records


@router.get(
    "/companies/{company_id}/compliance/summary",
    response_model=CompanyEssential8ComplianceSummary,
)
async def get_company_compliance_summary(
    company_id: int,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Get a summary of compliance status for a company"""
    # Super admins can view any company, regular users can only view their own
    is_super_admin = user.get("is_super_admin", False)
    user_company_id = user.get("company_id")
    
    if not is_super_admin and user_company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only view compliance for your own company",
        )
    
    summary = await essential8_repo.get_company_compliance_summary(company_id)
    return summary


@router.post(
    "/companies/{company_id}/compliance/initialize",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
)
async def initialize_company_compliance(
    company_id: int,
    _: None = Depends(require_database),
    __: dict = Depends(require_super_admin),
):
    """Initialize compliance records for all Essential 8 controls for a company"""
    created_count = await essential8_repo.initialize_company_compliance(company_id)
    return {
        "message": "Compliance records initialized",
        "created_count": created_count,
    }


@router.post(
    "/companies/{company_id}/compliance",
    response_model=CompanyEssential8ComplianceResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_company_compliance(
    company_id: int,
    payload: CompanyEssential8ComplianceCreate,
    _: None = Depends(require_database),
    __: dict = Depends(require_super_admin),
):
    """Create a compliance record for a company"""
    # Ensure the company_id in the payload matches the URL
    if payload.company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Company ID in payload does not match URL",
        )
    
    # Check if record already exists
    existing = await essential8_repo.get_company_compliance(
        company_id=company_id,
        control_id=payload.control_id,
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Compliance record already exists for this control",
        )
    
    record = await essential8_repo.create_company_compliance(
        **payload.model_dump(),
    )
    return record


@router.get(
    "/companies/{company_id}/compliance/{control_id}",
    response_model=CompanyEssential8ComplianceResponse,
)
async def get_company_compliance(
    company_id: int,
    control_id: int,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Get a specific compliance record for a company and control"""
    # Super admins can view any company, regular users can only view their own
    is_super_admin = user.get("is_super_admin", False)
    user_company_id = user.get("company_id")
    
    if not is_super_admin and user_company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only view compliance for your own company",
        )
    
    record = await essential8_repo.get_company_compliance(
        company_id=company_id,
        control_id=control_id,
    )
    if not record:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Compliance record not found",
        )
    return record


@router.patch(
    "/companies/{company_id}/compliance/{control_id}",
    response_model=CompanyEssential8ComplianceResponse,
)
async def update_company_compliance(
    company_id: int,
    control_id: int,
    payload: CompanyEssential8ComplianceUpdate,
    _: None = Depends(require_database),
    user: dict = Depends(require_super_admin),
):
    """Update a compliance record for a company"""
    # Check if record exists
    existing = await essential8_repo.get_company_compliance(
        company_id=company_id,
        control_id=control_id,
    )
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Compliance record not found",
        )
    
    # Update the record
    updates = payload.model_dump(exclude_unset=True)
    updated = await essential8_repo.update_company_compliance(
        company_id=company_id,
        control_id=control_id,
        user_id=user.get("id"),
        **updates,
    )
    
    if not updated:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to update compliance record",
        )
    
    return updated


@router.get(
    "/companies/{company_id}/compliance/{control_id}/audit",
    response_model=list[CompanyEssential8AuditResponse],
)
async def list_compliance_audit(
    company_id: int,
    control_id: int,
    limit: int = 100,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """List audit trail for compliance changes"""
    # Super admins can view any company, regular users can only view their own
    is_super_admin = user.get("is_super_admin", False)
    user_company_id = user.get("company_id")
    
    if not is_super_admin and user_company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only view audit trail for your own company",
        )
    
    audit = await essential8_repo.list_compliance_audit(
        company_id=company_id,
        control_id=control_id,
        limit=limit,
    )
    return audit


# =============================================================================
# Essential 8 Requirements Endpoints
# =============================================================================


@router.get("/requirements", response_model=list[Essential8RequirementResponse])
async def list_requirements(
    control_id: Optional[int] = None,
    maturity_level: Optional[str] = None,
    _: None = Depends(require_database),
    __: dict = Depends(get_current_user),
):
    """List all Essential 8 requirements"""
    requirements = await essential8_repo.list_essential8_requirements(
        control_id=control_id,
        maturity_level=maturity_level,
    )
    return requirements


@router.get(
    "/controls/{control_id}/with-requirements",
    response_model=Essential8ControlWithRequirementsResponse,
)
async def get_control_with_requirements(
    control_id: int,
    company_id: Optional[int] = None,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Get a control with all its requirements grouped by maturity level"""
    # If company_id is provided, check permissions
    if company_id:
        is_super_admin = user.get("is_super_admin", False)
        user_company_id = user.get("company_id")
        
        if not is_super_admin and user_company_id != company_id:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="You can only view compliance for your own company",
            )
    
    control_data = await essential8_repo.get_control_with_requirements(
        control_id=control_id,
        company_id=company_id,
    )
    
    if not control_data:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Control not found",
        )
    
    return control_data


@router.post(
    "/companies/{company_id}/requirements/initialize",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
)
async def initialize_company_requirement_compliance(
    company_id: int,
    control_id: Optional[int] = None,
    _: None = Depends(require_database),
    __: dict = Depends(require_super_admin),
):
    """Initialize requirement compliance records for a company"""
    created_count = await essential8_repo.initialize_company_requirement_compliance(
        company_id=company_id,
        control_id=control_id,
    )
    return {
        "message": "Requirement compliance records initialized",
        "created_count": created_count,
    }


@router.get(
    "/companies/{company_id}/requirements/compliance",
    response_model=list[CompanyEssential8RequirementComplianceResponse],
)
async def list_company_requirement_compliance(
    company_id: int,
    control_id: Optional[int] = None,
    maturity_level: Optional[str] = None,
    status_filter: Optional[ComplianceStatus] = None,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """List requirement compliance records for a company"""
    await _assert_company_compliance_access(user, company_id)
    records = await essential8_repo.list_company_requirement_compliance(
        company_id=company_id,
        control_id=control_id,
        maturity_level=maturity_level,
        status=status_filter,
    )
    return records


@router.post(
    "/companies/{company_id}/requirements/compliance",
    response_model=CompanyEssential8RequirementComplianceResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_company_requirement_compliance(
    company_id: int,
    payload: CompanyEssential8RequirementComplianceCreate,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Create a requirement compliance record"""
    await _assert_company_compliance_access(user, company_id, write=True)
    if payload.company_id != company_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Company ID in payload does not match URL",
        )
    
    # Check if record already exists
    existing = await essential8_repo.get_company_requirement_compliance(
        company_id=company_id,
        requirement_id=payload.requirement_id,
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Requirement compliance record already exists",
        )
    
    record = await essential8_repo.create_company_requirement_compliance(
        **payload.model_dump(),
    )
    return record


@router.get(
    "/companies/{company_id}/requirements/{requirement_id}/compliance",
    response_model=CompanyEssential8RequirementComplianceResponse,
)
async def get_company_requirement_compliance(
    company_id: int,
    requirement_id: int,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Get a specific requirement compliance record"""
    await _assert_company_compliance_access(user, company_id)
    record = await essential8_repo.get_company_requirement_compliance(
        company_id=company_id,
        requirement_id=requirement_id,
    )
    if not record:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement compliance record not found",
        )
    return record


@router.patch(
    "/companies/{company_id}/requirements/{requirement_id}/compliance",
    response_model=CompanyEssential8RequirementComplianceResponse,
)
async def update_company_requirement_compliance(
    company_id: int,
    requirement_id: int,
    payload: CompanyEssential8RequirementComplianceUpdate,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Update a requirement compliance record"""
    await _assert_company_compliance_access(user, company_id, write=True)
    # Check if record exists
    existing = await essential8_repo.get_company_requirement_compliance(
        company_id=company_id,
        requirement_id=requirement_id,
    )
    if not existing:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement compliance record not found",
        )
    
    # Get the requirement to find its control_id
    requirement = await essential8_repo.get_essential8_requirement(requirement_id)
    if not requirement:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Requirement not found",
        )
    
    # Update the record
    updates = payload.model_dump(exclude_unset=True)
    if "approval_status" in updates and "approved_at" not in updates:
        if updates["approval_status"] == ApprovalStatus.APPROVED:
            updates["approved_at"] = datetime.now(timezone.utc)
            updates["approved_by"] = user.get("id")
        elif updates["approval_status"] == ApprovalStatus.CHANGES_REQUESTED:
            updates["approved_at"] = None
            updates["approved_by"] = user.get("id")
    updated = await essential8_repo.update_company_requirement_compliance(
        company_id=company_id,
        requirement_id=requirement_id,
        **updates,
    )
    
    if not updated:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to update requirement compliance record",
        )
    
    await essential8_repo.append_requirement_audit(
        company_id=company_id,
        requirement_id=requirement_id,
        user_id=user.get("id"),
        action="update",
        from_status=existing.get("status"),
        to_status=updated.get("status"),
        approval_status=updated.get("approval_status"),
        change_summary="Requirement compliance updated.",
    )

    # Auto-update the control compliance based on requirement statuses
    await essential8_repo.auto_update_control_compliance_from_requirements(
        company_id=company_id,
        control_id=requirement["control_id"],
    )
    
    return updated


@router.post(
    "/companies/{company_id}/requirements/compliance/bulk",
    response_model=dict,
)
async def bulk_update_requirement_compliance(
    company_id: int,
    payload: CompanyEssential8RequirementBulkUpdate,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    """Apply the same requirement status/workflow update to multiple requirements."""
    await _assert_company_compliance_access(user, company_id, write=True)
    company_ids = payload.company_ids or [company_id]
    if any(int(candidate) != int(company_id) for candidate in company_ids) and not user.get("is_super_admin"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Company admins can only bulk update their active company",
        )
    requirement_ids = list(payload.requirement_ids)
    if payload.control_ids:
        requirements = await essential8_repo.list_essential8_requirements()
        requirement_ids.extend(
            int(requirement["id"])
            for requirement in requirements
            if int(requirement["control_id"]) in {int(control_id) for control_id in payload.control_ids}
        )
    requirement_ids = sorted({int(requirement_id) for requirement_id in requirement_ids})
    if not requirement_ids:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Select at least one requirement or control for the bulk update",
        )
    updates = payload.model_dump(exclude={"company_ids", "control_ids", "requirement_ids"}, exclude_unset=True)
    if updates.get("approval_status") == ApprovalStatus.APPROVED:
        updates.setdefault("approved_by", user.get("id"))
        updates.setdefault("approved_at", datetime.now(timezone.utc))
    try:
        return await essential8_repo.bulk_update_company_requirement_compliance(
            company_ids=company_ids,
            requirement_ids=requirement_ids,
            user_id=user.get("id"),
            **updates,
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.get(
    "/companies/{company_id}/requirements/{requirement_id}/evidence",
    response_model=list[Essential8RequirementEvidenceResponse],
)
async def list_requirement_evidence(
    company_id: int,
    requirement_id: int,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    await _assert_company_compliance_access(user, company_id)
    return await essential8_repo.list_requirement_evidence(company_id, requirement_id)


@router.post(
    "/companies/{company_id}/requirements/{requirement_id}/evidence/upload",
    response_model=Essential8RequirementEvidenceResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_requirement_evidence(
    company_id: int,
    requirement_id: int,
    title: str = Form(...),
    description: str | None = Form(default=None),
    evidence_file: UploadFile = File(...),
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    await _assert_company_compliance_access(user, company_id, write=True)
    requirement = await essential8_repo.get_essential8_requirement(requirement_id)
    if not requirement:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Requirement not found")
    safe_name = Path(evidence_file.filename or "evidence.bin").name
    if not safe_name:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Evidence file name is required")
    content = await evidence_file.read()
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    storage_name = f"company_{company_id}_requirement_{requirement_id}_{timestamp}_{safe_name}"
    storage_dir = _requirement_upload_dir()
    storage_path = storage_dir / storage_name
    storage_path.write_bytes(content)
    relative_path = f"compliance/essential8/{storage_name}"
    return await essential8_repo.add_requirement_evidence(
        company_id=company_id,
        requirement_id=requirement_id,
        title=title,
        description=description,
        file_name=safe_name,
        content_type=evidence_file.content_type,
        file_path=relative_path,
        file_size_bytes=len(content),
        uploaded_by=user.get("id"),
    )


@router.get(
    "/companies/{company_id}/requirements/reminders/summary",
    response_model=dict,
)
async def get_requirement_reminder_summary(
    company_id: int,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    await _assert_company_compliance_access(user, company_id)
    return await essential8_repo.get_requirement_reminder_summary(company_id)


@router.get(
    "/companies/{company_id}/requirements/trends",
    response_model=list[dict],
)
async def get_requirement_trends(
    company_id: int,
    control_id: Optional[int] = None,
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    await _assert_company_compliance_access(user, company_id)
    return await essential8_repo.get_requirement_trend(company_id, control_id=control_id)


@router.get(
    "/companies/{company_id}/requirements/export-bundle",
)
async def export_requirement_bundle(
    company_id: int,
    control_id: Optional[int] = None,
    export_format: str = Query("json", alias="format"),
    _: None = Depends(require_database),
    user: dict = Depends(get_current_user),
):
    await _assert_company_compliance_access(user, company_id)
    bundle = await essential8_repo.build_requirement_export_bundle(company_id, control_id=control_id)
    filename_root = f"essential8_audit_bundle_{company_id}"
    if export_format == "json":
        return bundle
    if export_format == "txt":
        lines = [f"Essential 8 compliance export for company {company_id}"]
        for record in bundle.get("requirements", []):
            lines.append(
                f"- {record.get('maturity_level')} requirement {record.get('requirement_order')}: "
                f"{record.get('status')} ({record.get('evidence_reference_count')} evidence refs)"
            )
        body = "\n".join(lines).encode("utf-8")
        return StreamingResponse(
            iter([body]),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename_root}.txt"'},
        )
    if export_format == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Essential 8 compliance audit bundle", 0)
        document.add_paragraph(f"Company ID: {company_id}")
        if control_id is not None:
            document.add_paragraph(f"Control ID filter: {control_id}")
        document.add_paragraph(f"Generated at: {bundle.get('generated_at')}")
        for record in bundle.get("requirements", []):
            heading = (
                f"{str(record.get('maturity_level') or '').upper()} requirement "
                f"{record.get('requirement_order') or '—'}"
            )
            document.add_heading(heading, level=1)
            document.add_paragraph(str(record.get("description") or ""))
            document.add_paragraph(
                f"Status: {record.get('status')}\n"
                f"Owner: {record.get('owner_user_id') or '—'}\n"
                f"Approval: {record.get('approval_status') or '—'}\n"
                f"Target date: {record.get('target_compliance_date') or '—'}"
            )
            document.add_paragraph(str(record.get("notes") or "No notes recorded."))
            if record.get("evidence_references"):
                document.add_paragraph("Evidence references:")
                for evidence in record["evidence_references"]:
                    document.add_paragraph(
                        f"v{evidence.get('version_number')} — {evidence.get('title')} "
                        f"({evidence.get('file_name')}) [{evidence.get('file_path')}]",
                        style="List Bullet",
                    )
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            iter([buffer.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename_root}.docx"'},
        )
    if export_format == "pdf":
        try:
            from weasyprint import HTML  # type: ignore
        except (ImportError, OSError) as exc:  # pragma: no cover
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="PDF export requires WeasyPrint and its native dependencies",
            ) from exc
        rows = []
        for record in bundle.get("requirements", []):
            evidence_html = "".join(
                (
                    "<li>"
                    f"v{html.escape(str(evidence.get('version_number') or ''))} — "
                    f"{html.escape(str(evidence.get('title') or ''))} "
                    f"({html.escape(str(evidence.get('file_name') or ''))}) "
                    f"[{html.escape(str(evidence.get('file_path') or ''))}]"
                    "</li>"
                )
                for evidence in record.get("evidence_references", [])
            ) or "<li>No evidence references</li>"
            rows.append(
                "<tr>"
                f"<td>{html.escape(str(record.get('control_id') or ''))}</td>"
                f"<td>{html.escape(str(record.get('maturity_level') or ''))}</td>"
                f"<td>{html.escape(str(record.get('requirement_order') or ''))}</td>"
                f"<td>{html.escape(str(record.get('description') or ''))}</td>"
                f"<td>{html.escape(str(record.get('status') or ''))}</td>"
                f"<td>{html.escape(str(record.get('approval_status') or ''))}</td>"
                f"<td><ul>{evidence_html}</ul></td>"
                "</tr>"
            )
        pdf_html = (
            "<html><body>"
            "<h1>Essential 8 compliance audit bundle</h1>"
            f"<p>Company ID: {company_id}</p>"
            f"<p>Generated at: {html.escape(str(bundle.get('generated_at') or ''))}</p>"
            "<table border='1' cellspacing='0' cellpadding='4'>"
            "<thead><tr><th>Control</th><th>Level</th><th>Req</th><th>Description</th>"
            "<th>Status</th><th>Approval</th><th>Evidence refs</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
            "</body></html>"
        )
        pdf_bytes = HTML(string=pdf_html).write_pdf()
        return StreamingResponse(
            iter([pdf_bytes]),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename_root}.pdf"'},
        )
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported export format")
